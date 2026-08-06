"""
Server-side audit report export (JSON, Markdown, CSV, HTML, XLSX, DOCX).

Optional sections: findings, architecture, hunt coverage, campaign summary,
mechanical codemap, PoC/evidence packs. Multi-format export can zip bundles.

needs_human = mech gates; confirmed = human accepted (not exploit proven).
Optional binary formats require optional deps (see pyproject [export] extra).
"""

from __future__ import annotations

import csv
import html
import io
import json
import zipfile
from dataclasses import dataclass, fields
from pathlib import Path
from typing import Any, Iterable

from vulnforge.util import utc_now_iso

DISCLAIMER = (
    "needs_human = mechanical gates passed; confirmed = human accepted "
    "(not exploit proven). Always re-check before production claims."
)

# Flat-row column order used by CSV/XLSX (and as a stable dict key set).
ROW_COLUMNS = [
    "id",
    "state",
    "severity",
    "weakness_class",
    "title",
    "summary",
    "path",
    "start_line",
    "end_line",
    "symbol",
    "attacker",
    "boundary",
    "impact",
    "evidence_id",
    "poc_relpath",
    "poc_files",
    "stable_key",
    "citations",
    "human_review_action",
    "human_review_notes",
    "validation_reasons",
    "poc_excerpt",  # filled when include_poc=True
]

# Caps when embedding evidence pack text into exports (PoC option).
POC_MAX_FILES = 12
POC_MAX_FILE_BYTES = 48_000
POC_MAX_EXCERPT_CHARS = 12_000

@dataclass
class ExportOptions:
    """What to include in an audit report export.

    Defaults keep pure-findings exports (tests / CLI) unless flags are set.
    The dashboard Export modal turns architecture + hunts + summary on by default.
    """

    include_findings: bool = True
    include_architecture: bool = False
    include_hunts: bool = False
    include_summary: bool = False
    include_poc: bool = False
    include_codemap: bool = False

    def as_dict(self) -> dict[str, bool]:
        return {f.name: bool(getattr(self, f.name)) for f in fields(self)}

    @classmethod
    def from_kwargs(
        cls,
        *,
        include_poc: bool = False,
        include_findings: bool = True,
        include_architecture: bool = False,
        include_hunts: bool = False,
        include_summary: bool = False,
        include_codemap: bool = False,
        options: ExportOptions | None = None,
    ) -> ExportOptions:
        if options is not None:
            return options
        return cls(
            include_findings=bool(include_findings),
            include_architecture=bool(include_architecture),
            include_hunts=bool(include_hunts),
            include_summary=bool(include_summary),
            include_poc=bool(include_poc),
            include_codemap=bool(include_codemap),
        )


def _run_meta(run_dir: Path | None, db: Any) -> dict[str, Any]:
    """Best-effort target/run identity for export headers."""
    meta: dict[str, Any] = {
        "run_dir": str(run_dir) if run_dir else "",
        "target_path": "",
        "run_id": "",
        "profile": "",
    }
    try:
        row = db.get_run() if db is not None else None
    except Exception:
        row = None
    if row is not None:
        # sqlite3.Row or mapping
        try:
            meta["target_path"] = row["target_path"] or ""
            meta["run_id"] = row["id"] or ""
            meta["profile"] = row["profile"] or ""
        except (KeyError, TypeError, IndexError):
            pass
    if run_dir is not None:
        p = Path(run_dir)
        meta.setdefault("run_id", p.name)
        if not meta["run_id"]:
            meta["run_id"] = p.name
    return meta


def _safe_arch(db: Any) -> dict[str, Any] | None:
    try:
        arch = db.get_architecture() if db is not None else None
    except Exception:
        return None
    return arch if isinstance(arch, dict) and arch else None


def _safe_codemap(db: Any) -> dict[str, Any] | None:
    try:
        cm = db.get_codemap() if db is not None else None
    except Exception:
        return None
    return cm if isinstance(cm, dict) and cm else None


def _safe_coverage(db: Any) -> dict[str, Any] | None:
    try:
        m = db.coverage_matrix() if db is not None else None
    except Exception:
        return None
    return m if isinstance(m, dict) else None


def _safe_summary(db: Any) -> dict[str, Any] | None:
    try:
        s = db.summary() if db is not None else None
    except Exception:
        return None
    return s if isinstance(s, dict) else None


def _hunt_task_rows(db: Any, *, limit: int = 400) -> list[dict[str, Any]]:
    """Compact hunt task list for 'what was hunted' report sections."""
    out: list[dict[str, Any]] = []
    if db is None:
        return out
    try:
        tasks = db.list_tasks(limit=max(1, int(limit)))
    except Exception:
        return out
    for t in tasks:
        kind = getattr(t, "kind", None) if not isinstance(t, dict) else t.get("kind")
        if str(kind or "") != "hunt":
            continue
        payload = (
            getattr(t, "payload", None) if not isinstance(t, dict) else t.get("payload")
        )
        if not isinstance(payload, dict):
            payload = {}
        area = str(payload.get("area") or "").strip()
        cls = str(
            payload.get("class") or payload.get("attack_class") or ""
        ).strip()
        state = getattr(t, "state", None) if not isinstance(t, dict) else t.get("state")
        tid = getattr(t, "id", None) if not isinstance(t, dict) else t.get("id")
        notes = str(payload.get("notes") or payload.get("operator_notes") or "").strip()
        path_hints = payload.get("path_hints") or payload.get("paths") or []
        if not isinstance(path_hints, list):
            path_hints = []
        out.append(
            {
                "task_id": tid,
                "state": state or "",
                "area": area,
                "class": cls,
                "path_hints": [str(p) for p in path_hints[:12] if p],
                "notes": notes[:400] if notes else "",
            }
        )
    return out


def _codemap_brief(codemap: dict[str, Any] | None) -> dict[str, Any] | None:
    """Trim mechanical codemap for export (avoid megabyte symbol dumps)."""
    if not isinstance(codemap, dict) or not codemap:
        return None
    summary = codemap.get("summary") if isinstance(codemap.get("summary"), dict) else {}
    mods = [m for m in (codemap.get("modules") or []) if isinstance(m, dict)]
    eps = [e for e in (codemap.get("entrypoints") or []) if isinstance(e, dict)]
    return {
        "source": codemap.get("source") or "mechanical",
        "target_kind": codemap.get("target_kind") or "",
        "generated_at": codemap.get("generated_at") or "",
        "summary": {
            "file_count": summary.get("file_count"),
            "module_count": summary.get("module_count", len(mods)),
            "symbol_count": summary.get("symbol_count"),
            "entrypoint_count": summary.get("entrypoint_count", len(eps)),
            "symbol_backend": summary.get("symbol_backend")
            or codemap.get("symbol_backend"),
            "languages": summary.get("languages") or {},
            "package_roots": (summary.get("package_roots") or [])[:20],
        },
        "modules": [
            {
                "path": m.get("path"),
                "kind": m.get("kind"),
                "file_count": m.get("file_count"),
                "symbol_count": m.get("symbol_count"),
                "signals": (m.get("signals") or [])[:8],
            }
            for m in mods[:80]
        ],
        "entrypoints": [
            {
                "path": e.get("path"),
                "kind": e.get("kind") or e.get("marker"),
            }
            for e in eps[:40]
        ],
    }


def collect_report_context(
    run_dir: Path | str | None,
    db: Any,
    options: ExportOptions,
) -> dict[str, Any]:
    """Gather all selected sections once for multi-format export."""
    run_dir_p = Path(run_dir) if run_dir else None
    meta = _run_meta(run_dir_p, db)
    rows: list[dict[str, Any]] = []
    if options.include_findings or options.include_poc:
        # Findings still loaded when only PoC requested with findings off? Skip.
        if options.include_findings:
            rows = attach_poc_packs(
                findings_to_rows(db),
                run_dir_p,
                include_poc=options.include_poc,
            )
    counts = _summary_counts(rows) if rows else _summary_counts([])
    # When findings omitted, still report DB totals if summary requested
    if not options.include_findings and options.include_summary and db is not None:
        try:
            for f in db.list_findings():
                st = getattr(f, "state", None) if not isinstance(f, dict) else f.get("state")
                counts["total"] += 1
                s = str(st or "").lower()
                if s == "confirmed":
                    counts["confirmed"] += 1
                elif s == "needs_human":
                    counts["needs_human"] += 1
                elif s == "candidate":
                    counts["candidate"] += 1
                elif s.startswith("rejected") or s == "superseded":
                    counts["rejected"] += 1
                else:
                    counts["other"] += 1
        except Exception:
            pass

    arch = _safe_arch(db) if options.include_architecture else None
    coverage = _safe_coverage(db) if options.include_hunts else None
    hunt_tasks = _hunt_task_rows(db) if options.include_hunts else []
    campaign = _safe_summary(db) if options.include_summary else None
    codemap = (
        _codemap_brief(_safe_codemap(db)) if options.include_codemap else None
    )

    return {
        "exported_at": utc_now_iso(),
        "disclaimer": DISCLAIMER,
        "options": options.as_dict(),
        "meta": meta,
        "run_dir": meta.get("run_dir") or (str(run_dir_p) if run_dir_p else ""),
        "target_path": meta.get("target_path") or "",
        "run_id": meta.get("run_id") or "",
        "profile": meta.get("profile") or "",
        "counts": counts,
        "rows": rows,
        "architecture": arch,
        "coverage": coverage,
        "hunt_tasks": hunt_tasks,
        "campaign": campaign,
        "codemap": codemap,
    }


def _md_architecture_section(arch: dict[str, Any] | None) -> list[str]:
    lines = ["## Architecture", ""]
    if not arch:
        lines.extend(
            [
                "_No architecture map stored yet (run recon / submit_architecture)._",
                "",
            ]
        )
        return lines
    summary = str(arch.get("summary") or "").strip()
    if summary:
        lines.extend([summary, ""])
    source = arch.get("source") or arch.get("_source") or ""
    if source:
        lines.append(f"_Source: `{source}`_")
        lines.append("")
    comps = [c for c in (arch.get("components") or []) if c]
    lines.extend([f"### Components ({len(comps)})", ""])
    if not comps:
        lines.append("_None recorded._")
        lines.append("")
    else:
        for c in comps[:80]:
            if isinstance(c, str):
                lines.append(f"- **{c}**")
                continue
            if not isinstance(c, dict):
                continue
            name = c.get("name") or c.get("id") or "?"
            role = c.get("role") or c.get("kind") or ""
            head = f"- **{name}**"
            if role:
                head += f" — {role}"
            lines.append(head)
            hints = c.get("path_hints") or c.get("paths") or []
            if isinstance(hints, list) and hints:
                lines.append(
                    "  - paths: "
                    + ", ".join(f"`{h}`" for h in hints[:12] if h)
                )
            note = c.get("notes") or c.get("description") or c.get("role_detail")
            if note:
                lines.append(f"  - {str(note)[:400]}")
        lines.append("")

    bounds = arch.get("trust_boundaries") or []
    if bounds:
        lines.extend(["### Trust boundaries", ""])
        for b in bounds[:40]:
            if isinstance(b, dict):
                label = b.get("name") or b.get("label") or b.get("boundary") or str(b)
                detail = b.get("description") or b.get("notes") or ""
                lines.append(f"- {label}" + (f" — {detail}" if detail else ""))
            else:
                lines.append(f"- {b}")
        lines.append("")

    surfaces = (
        arch.get("attack_surfaces")
        or arch.get("entrypoints")
        or arch.get("surfaces")
        or []
    )
    if surfaces:
        lines.extend(["### Attack surfaces / entrypoints", ""])
        for s in surfaces[:40]:
            if isinstance(s, dict):
                label = (
                    s.get("name")
                    or s.get("path")
                    or s.get("surface")
                    or s.get("kind")
                    or "?"
                )
                detail = s.get("description") or s.get("notes") or ""
                lines.append(f"- `{label}`" + (f" — {detail}" if detail else ""))
            else:
                lines.append(f"- {s}")
        lines.append("")

    focus = arch.get("hunt_focus") or arch.get("priority_areas") or []
    if focus:
        lines.extend(["### Hunt focus", ""])
        for item in focus[:30]:
            if isinstance(item, dict):
                label = item.get("area") or item.get("name") or item.get("class") or "?"
                why = item.get("why") or item.get("reason") or item.get("notes") or ""
                lines.append(f"- **{label}**" + (f" — {why}" if why else ""))
            else:
                lines.append(f"- {item}")
        lines.append("")

    flows = arch.get("data_flows") or arch.get("flows") or []
    if flows:
        lines.extend(["### Data flows (sample)", ""])
        for fl in flows[:20]:
            if isinstance(fl, dict):
                src = fl.get("from") or fl.get("source") or "?"
                dst = fl.get("to") or fl.get("sink") or "?"
                via = fl.get("via") or fl.get("channel") or ""
                bit = f"- `{src}` → `{dst}`"
                if via:
                    bit += f" via {via}"
                lines.append(bit)
            else:
                lines.append(f"- {fl}")
        lines.append("")
    return lines


def _md_coverage_section(
    coverage: dict[str, Any] | None, hunt_tasks: list[dict[str, Any]]
) -> list[str]:
    lines = [
        "## Hunts & coverage",
        "",
        "_Depth: empty/planned = not executed or residual, shallow = no substantial "
        "tools, none = honest miss, candidate/needs_human = filed. "
        "Shallow/aborted ≠ safe._",
        "",
    ]
    if coverage and (coverage.get("cells") or coverage.get("areas")):
        classes = coverage.get("classes") or []
        areas = coverage.get("areas") or []
        cell_map = {
            (c.get("area"), c.get("class")): c
            for c in (coverage.get("cells") or [])
            if isinstance(c, dict)
        }
        lines.extend(["### Coverage matrix (area × class)", ""])
        if classes and areas:
            # Cap columns for readability
            show_classes = classes[:16]
            extra = len(classes) - len(show_classes)
            header = "| area | " + " | ".join(show_classes) + " |"
            sep = "| --- | " + " | ".join(["---"] * len(show_classes)) + " |"
            lines.append(header)
            lines.append(sep)
            for a in areas[:60]:
                row = [str(a)]
                for cl in show_classes:
                    cell = cell_map.get((a, cl))
                    if not cell:
                        row.append(".")
                    else:
                        d = cell.get("last_depth") or "planned"
                        v = cell.get("visit_count") or 0
                        row.append(f"{d}×{v}" if v else str(d))
                lines.append("| " + " | ".join(row) + " |")
            if extra > 0:
                lines.append("")
                lines.append(f"_…{extra} more class column(s) omitted in table._")
            lines.append("")
        else:
            lines.append("_No matrix cells yet._")
            lines.append("")
    else:
        lines.append("_No coverage matrix yet (enqueue and run hunts)._")
        lines.append("")

    lines.extend([f"### Hunt tasks ({len(hunt_tasks)})", ""])
    if not hunt_tasks:
        lines.append("_No hunt tasks recorded._")
        lines.append("")
    else:
        # Group by state
        by_state: dict[str, int] = {}
        for h in hunt_tasks:
            st = str(h.get("state") or "unknown")
            by_state[st] = by_state.get(st, 0) + 1
        if by_state:
            lines.append(
                "Status: "
                + ", ".join(f"`{k}`={v}" for k, v in sorted(by_state.items()))
            )
            lines.append("")
        for h in hunt_tasks[:200]:
            area = h.get("area") or "?"
            cls = h.get("class") or "?"
            st = h.get("state") or "?"
            tid = h.get("task_id")
            bit = f"- **#{tid}** `{area}` × `{cls}` — {st}"
            hints = h.get("path_hints") or []
            if hints:
                bit += f" · paths: {', '.join(f'`{p}`' for p in hints[:4])}"
            lines.append(bit)
            if h.get("notes"):
                lines.append(f"  - notes: {h['notes']}")
        if len(hunt_tasks) > 200:
            lines.append(f"- _…{len(hunt_tasks) - 200} more tasks omitted._")
        lines.append("")
    return lines


def _md_summary_section(campaign: dict[str, Any] | None, counts: dict[str, int]) -> list[str]:
    lines = ["## Campaign summary", ""]
    if not campaign:
        lines.extend(
            [
                (
                    f"Findings — needs_human: {counts.get('needs_human', 0)}, "
                    f"confirmed: {counts.get('confirmed', 0)}, "
                    f"candidate: {counts.get('candidate', 0)}, "
                    f"rejected: {counts.get('rejected', 0)}, "
                    f"total: {counts.get('total', 0)}"
                ),
                "",
            ]
        )
        return lines
    run = campaign.get("run") if isinstance(campaign.get("run"), dict) else {}
    if run:
        lines.append(f"- **Target:** `{run.get('target_path') or '-'}`")
        lines.append(f"- **Profile:** `{run.get('profile') or '-'}`")
        if run.get("prompt_pin"):
            lines.append(f"- **Prompt pin:** `{run.get('prompt_pin')}`")
        if "has_architecture" in run:
            lines.append(
                f"- **Architecture:** "
                f"{'yes' if run.get('has_architecture') else 'no'}"
            )
        if "has_codemap" in run:
            lines.append(
                f"- **Codemap:** {'yes' if run.get('has_codemap') else 'no'}"
            )
    tasks = campaign.get("tasks")
    if isinstance(tasks, dict):
        lines.append(f"- **Tasks:** `{json.dumps(tasks, ensure_ascii=False)}`")
    findings = campaign.get("findings")
    if isinstance(findings, dict):
        lines.append(f"- **Findings by state:** `{json.dumps(findings, ensure_ascii=False)}`")
    else:
        lines.append(
            f"- **Findings counts:** needs_human={counts.get('needs_human', 0)}, "
            f"confirmed={counts.get('confirmed', 0)}, "
            f"candidate={counts.get('candidate', 0)}, "
            f"rejected={counts.get('rejected', 0)}, "
            f"total={counts.get('total', 0)}"
        )
    if campaign.get("has_work") is not None:
        lines.append(f"- **Has queued work:** {campaign.get('has_work')}")
    lines.append("")
    return lines


def _md_codemap_section(codemap: dict[str, Any] | None) -> list[str]:
    lines = ["## Mechanical codemap", ""]
    if not codemap:
        lines.extend(["_No mechanical codemap stored (run recon)._ ", ""])
        return lines
    s = codemap.get("summary") if isinstance(codemap.get("summary"), dict) else {}
    lines.extend(
        [
            f"- **Source:** {codemap.get('source') or 'mechanical'}",
            f"- **Files:** {s.get('file_count', '—')}",
            f"- **Modules:** {s.get('module_count', '—')}",
            f"- **Symbols:** {s.get('symbol_count', '—')}",
            f"- **Entrypoints:** {s.get('entrypoint_count', '—')}",
            "",
            "### Modules (capped)",
            "",
        ]
    )
    mods = codemap.get("modules") or []
    if not mods:
        lines.append("_None._")
    for m in mods[:60]:
        if not isinstance(m, dict):
            continue
        path_s = m.get("path") or "?"
        nfiles = m.get("file_count")
        nsym = m.get("symbol_count")
        extra = []
        if nfiles is not None:
            extra.append(f"{nfiles} files")
        if nsym is not None:
            extra.append(f"{nsym} symbols")
        lines.append(
            f"- `{path_s}`"
            + (f" ({', '.join(extra)})" if extra else "")
        )
    lines.append("")
    eps = codemap.get("entrypoints") or []
    if eps:
        lines.extend(["### Entrypoints", ""])
        for e in eps[:30]:
            if isinstance(e, dict):
                lines.append(f"- `{e.get('path')}` ({e.get('kind') or 'file'})")
        lines.append("")
    return lines


def _narrative_md_preamble(ctx: dict[str, Any]) -> list[str]:
    """Shared MD header + optional campaign/arch/hunts/codemap before findings."""
    opts = ctx.get("options") or {}
    lines = [
        "# VulnForge audit report",
        "",
        f"Target: `{ctx.get('target_path') or '-'}`",
        f"Run: `{ctx.get('run_id') or '-'}`",
        f"Profile: `{ctx.get('profile') or '-'}`",
        f"Run dir: `{ctx.get('run_dir') or '-'}`",
        f"Exported: {ctx.get('exported_at') or utc_now_iso()}",
        (
            "Sections: "
            + ", ".join(
                k.replace("include_", "")
                for k, v in sorted(opts.items())
                if v and str(k).startswith("include_")
            )
            or "none"
        ),
        f"Include PoC: `{'yes' if opts.get('include_poc') else 'no'}`",
        "",
        f"> **Disclaimer:** {DISCLAIMER}",
        "",
    ]
    counts = ctx.get("counts") or {}
    if opts.get("include_findings") or counts.get("total"):
        lines.extend(
            [
                (
                    f"Needs human: {counts.get('needs_human', 0)} | "
                    f"Confirmed (human): {counts.get('confirmed', 0)} | "
                    f"Candidates: {counts.get('candidate', 0)} | "
                    f"Rejected: {counts.get('rejected', 0)} | "
                    f"Total: {counts.get('total', 0)}"
                ),
                "",
            ]
        )
    if opts.get("include_summary"):
        lines.extend(_md_summary_section(ctx.get("campaign"), counts))
    if opts.get("include_architecture"):
        lines.extend(_md_architecture_section(ctx.get("architecture")))
    if opts.get("include_codemap"):
        lines.extend(_md_codemap_section(ctx.get("codemap")))
    if opts.get("include_hunts"):
        lines.extend(
            _md_coverage_section(ctx.get("coverage"), ctx.get("hunt_tasks") or [])
        )
    return lines


def _primary_citation(body: dict) -> dict[str, Any]:
    if body.get("sink_path"):
        return {
            "path": body.get("sink_path"),
            "start_line": None,
            "end_line": None,
            "symbol": body.get("sink_symbol"),
        }
    for c in body.get("citations") or []:
        if isinstance(c, dict) and c.get("path"):
            return c
    return {}


def _path_label(cit: dict | None) -> str:
    if not cit or not cit.get("path"):
        return ""
    s = str(cit["path"])
    start = cit.get("start_line")
    end = cit.get("end_line")
    if start is not None:
        s += f":{start}"
        if end is not None and end != start:
            s += f"-{end}"
    return s


def _severity(body: dict, finding: Any = None) -> str:
    if finding is not None and getattr(finding, "severity", None):
        return str(finding.severity)
    return str(body.get("severity_claim") or "unknown")


def _citations_flat(body: dict) -> str:
    parts = []
    for c in body.get("citations") or []:
        if not isinstance(c, dict) or not c.get("path"):
            continue
        label = _path_label(c)
        if c.get("symbol"):
            label = f"{label} ({c['symbol']})"
        parts.append(label)
    return "; ".join(parts)


def _human_review_latest(body: dict) -> dict[str, Any]:
    """Return latest human review entry if present."""
    latest = body.get("human_review_latest")
    if isinstance(latest, dict) and latest:
        return latest
    hist = body.get("human_review")
    if isinstance(hist, list) and hist:
        last = hist[-1]
        if isinstance(last, dict):
            return last
    return {}


def _validation_reasons_flat(body: dict) -> str:
    reasons = body.get("validation_reasons")
    if not isinstance(reasons, list) or not reasons:
        mech = body.get("validation_mech")
        if isinstance(mech, dict):
            reasons = mech.get("reasons") or []
    if not isinstance(reasons, list) or not reasons:
        return ""
    parts = []
    for r in reasons:
        if isinstance(r, str):
            parts.append(r)
        else:
            parts.append(json.dumps(r, ensure_ascii=False))
    return "; ".join(parts)


def load_evidence_pack(
    run_dir: Path | str | None,
    evidence_id: str | None,
    *,
    prefer_relpath: str | None = None,
    max_files: int = POC_MAX_FILES,
    max_file_bytes: int = POC_MAX_FILE_BYTES,
) -> dict[str, Any]:
    """
    Load text files from evidence/<id>/ for export (optional PoC inclusion).

    Returns {evidence_id, files: [{relpath, size, content, truncated}], preferred}.
    Never raises for missing packs — returns empty files list.
    """
    out: dict[str, Any] = {
        "evidence_id": evidence_id or "",
        "files": [],
        "preferred": prefer_relpath or "",
        "missing": False,
    }
    if not run_dir or not evidence_id:
        out["missing"] = True
        return out
    try:
        from vulnforge.tools.evidence_write import sanitize_evidence_id

        eid = sanitize_evidence_id(str(evidence_id))
    except Exception:
        out["missing"] = True
        return out
    root = (Path(run_dir) / "evidence" / eid).resolve()
    base = (Path(run_dir) / "evidence").resolve()
    try:
        root.relative_to(base)
    except ValueError:
        out["missing"] = True
        return out
    if not root.is_dir():
        out["missing"] = True
        return out

    files_meta: list[tuple[str, Path, int]] = []
    for f in sorted(root.rglob("*")):
        if not f.is_file() or f.name.endswith(".tmp"):
            continue
        try:
            rel = str(f.relative_to(root)).replace("\\", "/")
            size = f.stat().st_size
        except OSError:
            continue
        files_meta.append((rel, f, size))

    # Prefer declared PoC first, then .md, then rest
    def _sort_key(item: tuple[str, Path, int]) -> tuple[int, str]:
        rel = item[0]
        if prefer_relpath and rel == prefer_relpath:
            return (0, rel)
        if rel.lower().endswith(".md"):
            return (1, rel)
        return (2, rel)

    files_meta.sort(key=_sort_key)
    packed: list[dict[str, Any]] = []
    for rel, path, size in files_meta[: max(1, int(max_files))]:
        truncated = False
        try:
            raw = path.read_bytes()
        except OSError:
            continue
        if len(raw) > max_file_bytes:
            raw = raw[:max_file_bytes]
            truncated = True
        try:
            content = raw.decode("utf-8")
        except UnicodeDecodeError:
            content = raw.decode("utf-8", errors="replace")
            truncated = True
        if len(content) > POC_MAX_EXCERPT_CHARS:
            content = content[:POC_MAX_EXCERPT_CHARS] + "\n…[truncated]"
            truncated = True
        packed.append(
            {
                "relpath": rel,
                "size": size,
                "content": content,
                "truncated": truncated,
            }
        )
    out["files"] = packed
    if prefer_relpath and any(p["relpath"] == prefer_relpath for p in packed):
        out["preferred"] = prefer_relpath
    elif packed:
        out["preferred"] = packed[0]["relpath"]
    return out


def _preferred_poc_relpath(body: dict) -> str:
    poc = body.get("poc_relpath")
    return str(poc).strip() if poc else ""


def attach_poc_packs(
    rows: list[dict[str, Any]],
    run_dir: Path | str | None,
    *,
    include_poc: bool,
) -> list[dict[str, Any]]:
    """Optionally load evidence pack contents onto each finding row."""
    if not include_poc or not run_dir:
        for r in rows:
            r.setdefault("poc_files", "")
            r.setdefault("poc_excerpt", "")
            r.setdefault("poc_relpath", r.get("poc_relpath") or "")
        return rows
    cache: dict[str, dict[str, Any]] = {}
    for r in rows:
        body = r.get("_body") or {}
        eid = str(r.get("evidence_id") or body.get("evidence_id") or "")
        prefer = _preferred_poc_relpath(body) or str(r.get("poc_relpath") or "")
        r["poc_relpath"] = prefer
        if not eid:
            r["poc_files"] = ""
            r["poc_excerpt"] = ""
            r["_poc"] = {"evidence_id": "", "files": [], "missing": True}
            continue
        if eid not in cache:
            cache[eid] = load_evidence_pack(
                run_dir, eid, prefer_relpath=prefer or None
            )
        pack = cache[eid]
        # Re-prefer per finding if same pack has different poc_relpath
        if prefer and pack.get("preferred") != prefer:
            pack = load_evidence_pack(run_dir, eid, prefer_relpath=prefer)
        r["_poc"] = pack
        names = [f.get("relpath") or "" for f in pack.get("files") or []]
        r["poc_files"] = "; ".join(names)
        excerpt = ""
        for f in pack.get("files") or []:
            if prefer and f.get("relpath") == prefer:
                excerpt = f.get("content") or ""
                break
        if not excerpt and pack.get("files"):
            excerpt = pack["files"][0].get("content") or ""
        if len(excerpt) > 2000:
            excerpt = excerpt[:2000] + "…"
        r["poc_excerpt"] = excerpt
    return rows


def _partition_rows(rows: list[dict]) -> dict[str, list[dict]]:
    """Bucket rows for sectioned MD/HTML/DOCX reports."""
    buckets: dict[str, list[dict]] = {
        "needs_human": [],
        "confirmed": [],
        "candidate": [],
        "rejected": [],
        "other": [],
    }
    for r in rows:
        st = str(r.get("state") or "").lower()
        if st == "needs_human":
            buckets["needs_human"].append(r)
        elif st == "confirmed":
            buckets["confirmed"].append(r)
        elif st == "candidate":
            buckets["candidate"].append(r)
        elif st.startswith("rejected") or st == "superseded":
            buckets["rejected"].append(r)
        else:
            buckets["other"].append(r)
    return buckets


def findings_to_rows(db: Any) -> list[dict[str, Any]]:
    """Flatten all findings into exportable row dicts."""
    rows: list[dict[str, Any]] = []
    for f in db.list_findings():
        body = f.body if isinstance(getattr(f, "body", None), dict) else {}
        if not body and isinstance(f, dict):
            body = f.get("body") or {}
        tm = body.get("threat_model") or {}
        if not isinstance(tm, dict):
            tm = {}
        primary = _primary_citation(body)
        fid = getattr(f, "id", None) if not isinstance(f, dict) else f.get("id")
        state = getattr(f, "state", None) if not isinstance(f, dict) else f.get("state")
        stable = (
            getattr(f, "stable_key", None)
            if not isinstance(f, dict)
            else f.get("stable_key")
        )
        evid = (
            getattr(f, "evidence_id", None)
            if not isinstance(f, dict)
            else f.get("evidence_id")
        )
        evid = evid or body.get("evidence_id") or ""
        hr = _human_review_latest(body)
        rows.append(
            {
                "id": fid,
                "state": state or "",
                "severity": _severity(body, f if not isinstance(f, dict) else None),
                "weakness_class": body.get("weakness_class") or "",
                "title": body.get("title") or stable or f"Finding #{fid}",
                "summary": body.get("summary") or "",
                "path": primary.get("path") or "",
                "start_line": primary.get("start_line")
                if primary.get("start_line") is not None
                else "",
                "end_line": primary.get("end_line")
                if primary.get("end_line") is not None
                else "",
                "symbol": primary.get("symbol") or body.get("sink_symbol") or "",
                "attacker": tm.get("attacker") or "",
                "boundary": tm.get("boundary") or "",
                "impact": tm.get("impact") or "",
                "evidence_id": evid,
                "poc_relpath": _preferred_poc_relpath(body),
                "poc_files": "",
                "poc_excerpt": "",
                "stable_key": stable or "",
                "citations": _citations_flat(body),
                "human_review_action": hr.get("action") or "",
                "human_review_notes": (hr.get("notes") or "") if hr else "",
                "validation_reasons": _validation_reasons_flat(body),
                # Keep full body for richer formats
                "_body": body,
                "_finding": f,
                "_human_review": hr,
            }
        )
    return rows


def _summary_counts(rows: Iterable[dict]) -> dict[str, int]:
    c = {
        "confirmed": 0,
        "needs_human": 0,
        "candidate": 0,
        "rejected": 0,
        "other": 0,
        "total": 0,
    }
    for r in rows:
        c["total"] += 1
        st = str(r.get("state") or "").lower()
        if st == "confirmed":
            c["confirmed"] += 1
        elif st == "needs_human":
            c["needs_human"] += 1
        elif st == "candidate":
            c["candidate"] += 1
        elif st.startswith("rejected") or st == "superseded":
            c["rejected"] += 1
        else:
            c["other"] += 1
    return c


def _public_rows(rows: list[dict]) -> list[dict]:
    """Strip private keys used only during export composition."""
    out = []
    for r in rows:
        out.append({k: r[k] for k in ROW_COLUMNS if k in r})
    return out


def export_json(
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> str:
    """JSON document (UTF-8 text). Optional architecture / hunts / summary sections."""
    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    ctx = collect_report_context(run_dir, db, opts)
    rows = ctx.get("rows") or []
    findings_full = []
    if opts.include_findings:
        for r in rows:
            f = r.get("_finding")
            entry: dict[str, Any] = {
                "id": r.get("id"),
                "stable_key": r.get("stable_key"),
                "state": r.get("state"),
                "evidence_id": r.get("evidence_id"),
                "body": r.get("_body")
                or (getattr(f, "body", None) if f is not None else {}),
            }
            if opts.include_poc:
                entry["poc"] = r.get("_poc") or {
                    "evidence_id": r.get("evidence_id") or "",
                    "files": [],
                }
            findings_full.append(entry)
    payload: dict[str, Any] = {
        "exported_at": ctx.get("exported_at") or utc_now_iso(),
        "disclaimer": DISCLAIMER,
        "include_poc": bool(opts.include_poc),
        "sections": opts.as_dict(),
        "run_dir": ctx.get("run_dir") or "",
        "target_path": ctx.get("target_path") or "",
        "run_id": ctx.get("run_id") or "",
        "profile": ctx.get("profile") or "",
        "counts": ctx.get("counts") or _summary_counts(rows),
        "rows": _public_rows(rows) if opts.include_findings else [],
        "findings": findings_full,
    }
    if opts.include_architecture:
        payload["architecture"] = ctx.get("architecture")
    if opts.include_hunts:
        payload["coverage"] = ctx.get("coverage")
        payload["hunt_tasks"] = ctx.get("hunt_tasks") or []
    if opts.include_summary:
        payload["campaign"] = ctx.get("campaign")
    if opts.include_codemap:
        payload["codemap"] = ctx.get("codemap")
    return json.dumps(payload, indent=2, ensure_ascii=False) + "\n"

def _md_finding_block(r: dict) -> list[str]:
    """Markdown bullets + body for one finding row."""
    body = r.get("_body") or {}
    hr = r.get("_human_review") or {}
    lines = [
        f"### [{r['id']}] {r['title']}",
        "",
        f"- **state:** {r['state']}",
        f"- **class:** {r['weakness_class'] or '-'}",
        f"- **severity:** {r['severity'] or '-'}",
        f"- **stable_key:** `{r['stable_key'] or '-'}`",
    ]
    if r.get("attacker"):
        lines.append(f"- **attacker:** {r['attacker']}")
    if r.get("boundary"):
        lines.append(f"- **boundary:** {r['boundary']}")
    if r.get("impact"):
        lines.append(f"- **impact:** {r['impact']}")
    if r.get("path"):
        loc = _path_label(
            {
                "path": r["path"],
                "start_line": r["start_line"] if r["start_line"] != "" else None,
                "end_line": r["end_line"] if r["end_line"] != "" else None,
            }
        )
        lines.append(f"- **location:** `{loc}`")
    if r.get("evidence_id"):
        lines.append(f"- **evidence:** `{r['evidence_id']}`")
    if r.get("poc_relpath"):
        lines.append(f"- **poc_relpath:** `{r['poc_relpath']}`")
    if hr.get("action"):
        notes = hr.get("notes") or ""
        note_bit = f" — {notes}" if notes else ""
        lines.append(f"- **human_review:** {hr.get('action')}{note_bit}")
    if r.get("validation_reasons"):
        lines.append(f"- **validation_reasons:** {r['validation_reasons']}")
    lines.append("")
    if r.get("summary"):
        lines.append(r["summary"])
        lines.append("")
    lines.append("**Citations:**")
    lines.append("")
    cites = body.get("citations") or []
    if cites:
        for cit in cites:
            if not isinstance(cit, dict) or not cit.get("path"):
                continue
            label = _path_label(cit)
            if cit.get("symbol"):
                lines.append(f"- `{label}` ({cit['symbol']})")
            else:
                lines.append(f"- `{label}`")
    else:
        lines.append("- _none_")
    lines.append("")
    # Optional PoC / evidence pack contents
    pack = r.get("_poc")
    if isinstance(pack, dict) and pack.get("files"):
        lines.append("**PoC / evidence pack:**")
        lines.append("")
        for f in pack["files"]:
            rel = f.get("relpath") or "?"
            trunc = " (truncated)" if f.get("truncated") else ""
            lines.append(f"#### `{rel}`{trunc}")
            lines.append("")
            lang = "markdown" if str(rel).lower().endswith(".md") else "text"
            lines.append(f"```{lang}")
            lines.append(str(f.get("content") or "").rstrip())
            lines.append("```")
            lines.append("")
    elif isinstance(pack, dict) and pack.get("missing") and r.get("evidence_id"):
        lines.append("_Evidence pack missing or empty on disk._")
        lines.append("")
    return lines


def export_markdown(
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> str:
    """Markdown audit report (UTF-8 text), optional architecture / hunts / findings."""
    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    ctx = collect_report_context(run_dir, db, opts)
    lines = _narrative_md_preamble(ctx)
    rows = ctx.get("rows") or []

    if opts.include_findings:
        lines.append("## Findings")
        lines.append("")
        if not rows:
            lines.append("_No findings._")
            lines.append("")
        else:
            buckets = _partition_rows(rows)
            sections = [
                ("needs_human", "### Needs human review (mech-passed)"),
                ("confirmed", "### Confirmed findings (human-accepted)"),
                ("candidate", "### Candidates"),
                ("rejected", "### Rejected"),
                ("other", "### Other"),
            ]
            for key, heading in sections:
                group = buckets.get(key) or []
                if key == "other" and not group:
                    continue
                lines.append(heading)
                lines.append("")
                if not group:
                    lines.append("_None._")
                    lines.append("")
                    continue
                for r in group:
                    lines.extend(_md_finding_block(r))
    return "\n".join(lines)

def export_csv(
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> str:
    """CSV (UTF-8 text with header). Flat findings only; other sections ignored."""
    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    # CSV is findings-oriented; force findings on when this format is chosen alone.
    if not opts.include_findings:
        opts = ExportOptions(
            include_findings=True,
            include_architecture=False,
            include_hunts=False,
            include_summary=False,
            include_poc=opts.include_poc,
            include_codemap=False,
        )
    ctx = collect_report_context(run_dir, db, opts)
    rows = _public_rows(ctx.get("rows") or [])
    cols = list(ROW_COLUMNS)
    if not opts.include_poc:
        cols = [c for c in cols if c != "poc_excerpt"]
    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=cols,
        extrasaction="ignore",
        lineterminator="\n",
    )
    writer.writeheader()
    for r in rows:
        writer.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in cols})
    body = buf.getvalue()
    note = f"# {DISCLAIMER}"
    if opts.include_poc:
        note += " | include_poc=true (poc_excerpt truncated)"
    note += " | architecture/hunts/summary not in CSV (use md/html/json/docx/xlsx)"
    return f"{note}\n{body}"

def _esc(s: Any) -> str:
    return html.escape(str(s if s is not None else ""), quote=True)


def _html_from_md_lines(md_lines: list[str]) -> str:
    """Very small markdown→HTML for pre-built section lines (headings, lists, paragraphs)."""
    parts: list[str] = []
    in_ul = False
    for raw in md_lines:
        line = raw.rstrip()
        if not line:
            if in_ul:
                parts.append("</ul>")
                in_ul = False
            continue
        if line.startswith("### "):
            if in_ul:
                parts.append("</ul>")
                in_ul = False
            parts.append(f"<h3>{_esc(line[4:])}</h3>")
            continue
        if line.startswith("## "):
            if in_ul:
                parts.append("</ul>")
                in_ul = False
            parts.append(f"<h2>{_esc(line[3:])}</h2>")
            continue
        if line.startswith("- "):
            if not in_ul:
                parts.append("<ul>")
                in_ul = True
            item = line[2:]
            out: list[str] = []
            bits = item.split("`")
            for i, b in enumerate(bits):
                if i % 2 == 1:
                    out.append(f"<code>{_esc(b)}</code>")
                else:
                    out.append(_esc(b))
            parts.append(f"<li>{''.join(out)}</li>")
            continue
        if line.startswith(">"):
            if in_ul:
                parts.append("</ul>")
                in_ul = False
            parts.append(f"<blockquote>{_esc(line.lstrip('> ').strip())}</blockquote>")
            continue
        if line.startswith("| ") and "|" in line[2:]:
            if in_ul:
                parts.append("</ul>")
                in_ul = False
            # pass table lines as pre mono block accum — handled by caller for coverage
            parts.append(f"<pre class='table-line'>{_esc(line)}</pre>")
            continue
        if in_ul:
            parts.append("</ul>")
            in_ul = False
        if line.startswith("_") and line.endswith("_") and len(line) > 2:
            parts.append(f"<p class='muted'><em>{_esc(line[1:-1])}</em></p>")
        else:
            parts.append(f"<p>{_esc(line)}</p>")
    if in_ul:
        parts.append("</ul>")
    return "\n".join(parts)


def export_html(
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> str:
    """Self-contained HTML report with optional architecture/hunts + findings."""
    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    ctx = collect_report_context(run_dir, db, opts)
    rows = ctx.get("rows") or []
    counts = ctx.get("counts") or _summary_counts(rows)
    exported = ctx.get("exported_at") or utc_now_iso()
    meta = {
        "target_path": ctx.get("target_path"),
        "run_id": ctx.get("run_id"),
        "run_dir": ctx.get("run_dir"),
        "profile": ctx.get("profile"),
    }

    def _detail_block(r: dict) -> str:
        body = r.get("_body") or {}
        fid = r["id"]
        state_cls = str(r["state"] or "").lower().replace("_", "-")
        cite_lis = []
        for cit in body.get("citations") or []:
            if not isinstance(cit, dict) or not cit.get("path"):
                continue
            label = _path_label(cit)
            if cit.get("symbol"):
                cite_lis.append(f"<li><code>{_esc(label)}</code> ({_esc(cit['symbol'])})</li>")
            else:
                cite_lis.append(f"<li><code>{_esc(label)}</code></li>")
        cites_html = (
            f"<ul>{''.join(cite_lis)}</ul>" if cite_lis else "<p class='muted'>None</p>"
        )
        hr = r.get("_human_review") or {}
        hr_html = ""
        if hr.get("action"):
            hr_html = (
                f"<dt>Human review</dt><dd>{_esc(hr.get('action') or '-')}"
                f"{(' — ' + _esc(hr['notes'])) if hr.get('notes') else ''}</dd>"
            )
        vr = r.get("validation_reasons") or ""
        vr_html = (
            f"<dt>Validation reasons</dt><dd class='mono'>{_esc(vr)}</dd>" if vr else ""
        )
        poc_html = ""
        pack = r.get("_poc") if opts.include_poc else None
        if isinstance(pack, dict) and pack.get("files"):
            parts = []
            for f in pack["files"]:
                rel = f.get("relpath") or "?"
                trunc = " <em>(truncated)</em>" if f.get("truncated") else ""
                parts.append(
                    f"<h5><code>{_esc(rel)}</code>{trunc}</h5>"
                    f"<pre class='poc'>{_esc(f.get('content') or '')}</pre>"
                )
            poc_html = "<h4>PoC / evidence pack</h4>" + "".join(parts)
        elif opts.include_poc and r.get("evidence_id"):
            poc_html = "<p class='muted'>Evidence pack missing or empty on disk.</p>"
        return (
            f'<details class="finding" id="f-{_esc(fid)}">'
            f"<summary><strong>[{_esc(fid)}]</strong> {_esc(r['title'])} "
            f'<span class="badge state-{_esc(state_cls)}">{_esc(r["state"])}</span></summary>'
            f'<div class="finding-body">'
            f"<p>{_esc(r.get('summary') or 'No summary.')}</p>"
            f"<dl class='kv'>"
            f"<dt>Class</dt><dd class='mono'>{_esc(r['weakness_class'] or '-')}</dd>"
            f"<dt>Severity</dt><dd>{_esc(r['severity'] or '-')}</dd>"
            f"<dt>stable_key</dt><dd class='mono'>{_esc(r['stable_key'] or '-')}</dd>"
            f"<dt>Evidence</dt><dd class='mono'>{_esc(r['evidence_id'] or '-')}</dd>"
            f"<dt>PoC path</dt><dd class='mono'>{_esc(r.get('poc_relpath') or '-')}</dd>"
            f"<dt>Attacker</dt><dd>{_esc(r.get('attacker') or '-')}</dd>"
            f"<dt>Boundary</dt><dd>{_esc(r.get('boundary') or '-')}</dd>"
            f"<dt>Impact</dt><dd>{_esc(r.get('impact') or '-')}</dd>"
            f"{hr_html}"
            f"{vr_html}"
            f"</dl>"
            f"<h4>Citations</h4>{cites_html}"
            f"{poc_html}"
            f"</div></details>"
        )

    table_rows = []
    for r in rows:
        fid = r["id"]
        loc = _path_label(
            {
                "path": r["path"],
                "start_line": r["start_line"] if r["start_line"] != "" else None,
                "end_line": r["end_line"] if r["end_line"] != "" else None,
            }
        )
        state_cls = str(r["state"] or "").lower().replace("_", "-")
        table_rows.append(
            "<tr>"
            f'<td class="mono">{_esc(fid)}</td>'
            f"<td>{_esc(r['title'])}</td>"
            f'<td class="mono">{_esc(r["weakness_class"])}</td>'
            f"<td>{_esc(r['severity'])}</td>"
            f'<td><span class="badge state-{_esc(state_cls)}">{_esc(r["state"])}</span></td>'
            f'<td class="mono">{_esc(loc or "-")}</td>'
            f'<td><a href="#f-{_esc(fid)}">Details</a></td>'
            "</tr>"
        )

    table_body = (
        "\n".join(table_rows)
        if table_rows
        else '<tr><td colspan="7" class="empty">No findings.</td></tr>'
    )

    buckets = _partition_rows(rows)
    detail_sections = [
        ("needs_human", "Needs human review (mech-passed)"),
        ("confirmed", "Confirmed findings (human-accepted)"),
        ("candidate", "Candidates"),
        ("rejected", "Rejected"),
        ("other", "Other"),
    ]
    detail_parts: list[str] = []
    if opts.include_findings:
        if not rows:
            detail_parts.append("<p class='muted'>No findings.</p>")
        else:
            for key, heading in detail_sections:
                group = buckets.get(key) or []
                if key == "other" and not group:
                    continue
                detail_parts.append(f"<h3>{_esc(heading)}</h3>")
                if not group:
                    detail_parts.append("<p class='muted'>None.</p>")
                else:
                    detail_parts.extend(_detail_block(r) for r in group)
    details_html = "\n".join(detail_parts)

    # Narrative sections (summary / architecture / hunts / codemap) as HTML
    extra_md: list[str] = []
    if opts.include_summary:
        extra_md.extend(_md_summary_section(ctx.get("campaign"), counts))
    if opts.include_architecture:
        extra_md.extend(_md_architecture_section(ctx.get("architecture")))
    if opts.include_codemap:
        extra_md.extend(_md_codemap_section(ctx.get("codemap")))
    if opts.include_hunts:
        extra_md.extend(
            _md_coverage_section(ctx.get("coverage"), ctx.get("hunt_tasks") or [])
        )
    extra_html = _html_from_md_lines(extra_md) if extra_md else ""

    sections_label = ", ".join(
        k.replace("include_", "")
        for k, v in sorted(opts.as_dict().items())
        if v and k.startswith("include_")
    ) or "none"

    findings_section = ""
    if opts.include_findings:
        findings_section = f"""
  <section>
    <h2>Findings table</h2>
    <table>
      <thead>
        <tr>
          <th>ID</th><th>Title</th><th>Class</th><th>Severity</th>
          <th>State</th><th>Location</th><th></th>
        </tr>
      </thead>
      <tbody>
{table_body}
      </tbody>
    </table>
  </section>
  <section>
    <h2>Expandable detail</h2>
{details_html}
  </section>"""

    # ASCII-safe CSS and structure; no external assets.
    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>VulnForge audit report</title>
<style>
  :root {{
    --bg: #0f1419;
    --card: #1a2332;
    --text: #e7ecf3;
    --muted: #8b9bb4;
    --border: #2a3a4f;
    --accent: #3d8bfd;
    --good: #3dd68c;
    --warn: #f5a524;
    --bad: #f31260;
  }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    font-family: system-ui, -apple-system, Segoe UI, Roboto, Helvetica, Arial, sans-serif;
    background: var(--bg);
    color: var(--text);
    line-height: 1.45;
    padding: 1.5rem;
  }}
  h1, h2, h3, h4 {{ margin: 0 0 0.6rem; font-weight: 600; }}
  .muted {{ color: var(--muted); }}
  .disclaimer {{
    border-left: 3px solid var(--warn);
    background: var(--card);
    padding: 0.75rem 1rem;
    margin: 1rem 0 1.5rem;
  }}
  .stats {{
    display: flex;
    flex-wrap: wrap;
    gap: 0.75rem;
    margin-bottom: 1.25rem;
  }}
  .stat {{
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 8px;
    padding: 0.75rem 1rem;
    min-width: 7rem;
  }}
  .stat .label {{ color: var(--muted); font-size: 0.8rem; text-transform: uppercase; }}
  .stat .value {{ font-size: 1.4rem; font-weight: 600; }}
  .meta {{ color: var(--muted); font-size: 0.9rem; margin-bottom: 1rem; }}
  .meta code {{ color: var(--text); }}
  table {{
    width: 100%;
    border-collapse: collapse;
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 8px;
    overflow: hidden;
  }}
  th, td {{
    text-align: left;
    padding: 0.55rem 0.75rem;
    border-bottom: 1px solid var(--border);
    vertical-align: top;
  }}
  th {{ color: var(--muted); font-size: 0.8rem; text-transform: uppercase; letter-spacing: 0.03em; }}
  tr:last-child td {{ border-bottom: none; }}
  a {{ color: var(--accent); }}
  .mono {{ font-family: ui-monospace, SFMono-Regular, Menlo, Consolas, monospace; font-size: 0.9em; }}
  pre.poc, pre.table-line {{
    background: #0d1117;
    border: 1px solid var(--border);
    border-radius: 8px;
    padding: 0.75rem 1rem;
    overflow-x: auto;
    white-space: pre-wrap;
    word-break: break-word;
    font-size: 0.82rem;
    max-height: 28rem;
  }}
  pre.table-line {{ max-height: none; margin: 0.15rem 0; padding: 0.25rem 0.5rem; }}
  .badge {{
    display: inline-block;
    padding: 0.1rem 0.45rem;
    border-radius: 999px;
    font-size: 0.75rem;
    background: #243044;
    border: 1px solid var(--border);
  }}
  .state-confirmed {{ border-color: var(--good); color: var(--good); }}
  .state-needs-human {{ border-color: var(--warn); color: var(--warn); }}
  .state-candidate {{ border-color: var(--warn); color: var(--warn); }}
  .state-rejected-mech, .state-rejected-llm, .state-rejected-human, .state-superseded {{
    border-color: var(--bad); color: var(--bad);
  }}
  .empty {{ color: var(--muted); text-align: center; padding: 1.5rem; }}
  details.finding {{
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 8px;
    margin: 0.6rem 0;
    padding: 0.5rem 0.9rem;
  }}
  details.finding summary {{ cursor: pointer; padding: 0.35rem 0; }}
  .finding-body {{ padding: 0.5rem 0 0.75rem; }}
  dl.kv {{
    display: grid;
    grid-template-columns: 8rem 1fr;
    gap: 0.35rem 0.75rem;
    margin: 0.75rem 0;
  }}
  dl.kv dt {{ color: var(--muted); }}
  dl.kv dd {{ margin: 0; }}
  section {{ margin-top: 2rem; }}
  section.narrative ul {{ margin: 0.4rem 0 0.8rem 1.1rem; }}
  section.narrative li {{ margin: 0.2rem 0; }}
  footer {{ margin-top: 2rem; color: var(--muted); font-size: 0.85rem; }}
</style>
</head>
<body>
  <h1>VulnForge audit report</h1>
  <div class="meta">
    Target: <code>{_esc(meta.get("target_path") or "-")}</code><br>
    Run: <code>{_esc(meta.get("run_id") or "-")}</code><br>
    Profile: <code>{_esc(meta.get("profile") or "-")}</code><br>
    Run dir: <code>{_esc(meta.get("run_dir") or "-")}</code><br>
    Exported: {_esc(exported)}<br>
    Sections: <code>{_esc(sections_label)}</code><br>
    Include PoC: <code>{"yes" if opts.include_poc else "no"}</code>
  </div>
  <div class="disclaimer"><strong>Disclaimer:</strong> {_esc(DISCLAIMER)}</div>
  <div class="stats">
    <div class="stat"><div class="label">Needs human</div><div class="value">{counts.get("needs_human", 0)}</div></div>
    <div class="stat"><div class="label">Confirmed (human)</div><div class="value">{counts.get("confirmed", 0)}</div></div>
    <div class="stat"><div class="label">Candidates</div><div class="value">{counts.get("candidate", 0)}</div></div>
    <div class="stat"><div class="label">Rejected</div><div class="value">{counts.get("rejected", 0)}</div></div>
    <div class="stat"><div class="label">Total</div><div class="value">{counts.get("total", 0)}</div></div>
  </div>
  <section class="narrative">
{extra_html}
  </section>
{findings_section}
  <footer>Generated by VulnForge. harness.db is authoritative; this file is a snapshot export.</footer>
</body>
</html>
"""

def export_xlsx(
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> bytes:
    """Excel workbook (.xlsx). Requires openpyxl. Extra sheets for arch/hunts."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as e:
        raise ImportError(
            "export_xlsx requires openpyxl. Install with: pip install 'vulnforge[export]' "
            "or: pip install openpyxl>=3.1"
        ) from e

    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    ctx = collect_report_context(run_dir, db, opts)
    full_rows = ctx.get("rows") or []
    rows = _public_rows(full_rows)
    counts = ctx.get("counts") or _summary_counts(rows)
    meta = {
        "target_path": ctx.get("target_path"),
        "run_id": ctx.get("run_id"),
        "run_dir": ctx.get("run_dir"),
        "profile": ctx.get("profile"),
    }
    cols = list(ROW_COLUMNS)
    if not opts.include_poc:
        cols = [c for c in cols if c != "poc_excerpt"]

    wb = Workbook()
    # Meta first
    meta_ws = wb.active
    meta_ws.title = "Meta"
    meta_ws["A1"] = "VulnForge audit report"
    meta_ws["A1"].font = Font(bold=True, size=14)
    meta_ws["A3"] = "Disclaimer"
    meta_ws["B3"] = DISCLAIMER
    meta_ws["A4"] = "Exported"
    meta_ws["B4"] = ctx.get("exported_at") or utc_now_iso()
    meta_ws["A5"] = "Target"
    meta_ws["B5"] = meta.get("target_path") or ""
    meta_ws["A6"] = "Run id"
    meta_ws["B6"] = meta.get("run_id") or ""
    meta_ws["A7"] = "Profile"
    meta_ws["B7"] = meta.get("profile") or ""
    meta_ws["A8"] = "Run dir"
    meta_ws["B8"] = meta.get("run_dir") or ""
    meta_ws["A9"] = "Include PoC"
    meta_ws["B9"] = "yes" if opts.include_poc else "no"
    meta_ws["A10"] = "Sections"
    meta_ws["B10"] = ", ".join(
        k.replace("include_", "")
        for k, v in sorted(opts.as_dict().items())
        if v and k.startswith("include_")
    )
    meta_ws["A12"] = "Needs human"
    meta_ws["B12"] = counts.get("needs_human", 0)
    meta_ws["A13"] = "Confirmed (human)"
    meta_ws["B13"] = counts.get("confirmed", 0)
    meta_ws["A14"] = "Candidates"
    meta_ws["B14"] = counts.get("candidate", 0)
    meta_ws["A15"] = "Rejected"
    meta_ws["B15"] = counts.get("rejected", 0)
    meta_ws["A16"] = "Total"
    meta_ws["B16"] = counts.get("total", 0)
    meta_ws.column_dimensions["A"].width = 18
    meta_ws.column_dimensions["B"].width = 80

    if opts.include_findings:
        ws = wb.create_sheet("Findings")
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(val), 60))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[col_letter].width = max(10, max_len + 2)

    if opts.include_poc and opts.include_findings:
        poc_ws = wb.create_sheet("PoC packs")
        poc_ws.append(
            ["finding_id", "evidence_id", "relpath", "size", "truncated", "content"]
        )
        for cell in poc_ws[1]:
            cell.font = Font(bold=True)
        for r in full_rows:
            pack = r.get("_poc") or {}
            for f in pack.get("files") or []:
                poc_ws.append(
                    [
                        r.get("id"),
                        r.get("evidence_id"),
                        f.get("relpath"),
                        f.get("size"),
                        bool(f.get("truncated")),
                        f.get("content") or "",
                    ]
                )
        for col in poc_ws.columns:
            col_letter = col[0].column_letter
            poc_ws.column_dimensions[col_letter].width = 18 if col_letter != "F" else 60
            for cell in col:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    if opts.include_architecture:
        arch = ctx.get("architecture") or {}
        arch_ws = wb.create_sheet("Architecture")
        arch_ws["A1"] = "Summary"
        arch_ws["B1"] = str(arch.get("summary") or "") if arch else ""
        arch_ws["A2"] = "Source"
        arch_ws["B2"] = str(arch.get("source") or arch.get("_source") or "") if arch else ""
        arch_ws.append([])
        arch_ws.append(["component", "role", "path_hints", "notes"])
        for cell in arch_ws[4]:
            cell.font = Font(bold=True)
        for c in (arch.get("components") or []) if arch else []:
            if isinstance(c, str):
                arch_ws.append([c, "", "", ""])
            elif isinstance(c, dict):
                hints = c.get("path_hints") or c.get("paths") or []
                if isinstance(hints, list):
                    hints_s = "; ".join(str(h) for h in hints[:12])
                else:
                    hints_s = str(hints)
                arch_ws.append(
                    [
                        c.get("name") or c.get("id") or "",
                        c.get("role") or c.get("kind") or "",
                        hints_s,
                        str(c.get("notes") or c.get("description") or "")[:500],
                    ]
                )
        arch_ws.column_dimensions["A"].width = 24
        arch_ws.column_dimensions["B"].width = 20
        arch_ws.column_dimensions["C"].width = 40
        arch_ws.column_dimensions["D"].width = 50

    if opts.include_hunts:
        hunts_ws = wb.create_sheet("Hunt tasks")
        hunts_ws.append(["task_id", "state", "area", "class", "path_hints", "notes"])
        for cell in hunts_ws[1]:
            cell.font = Font(bold=True)
        for h in ctx.get("hunt_tasks") or []:
            hunts_ws.append(
                [
                    h.get("task_id"),
                    h.get("state"),
                    h.get("area"),
                    h.get("class"),
                    "; ".join(h.get("path_hints") or []),
                    h.get("notes") or "",
                ]
            )
        cov = ctx.get("coverage") or {}
        cov_ws = wb.create_sheet("Coverage")
        classes = cov.get("classes") or []
        areas = cov.get("areas") or []
        cell_map = {
            (c.get("area"), c.get("class")): c
            for c in (cov.get("cells") or [])
            if isinstance(c, dict)
        }
        cov_ws.append(["area"] + list(classes))
        for cell in cov_ws[1]:
            cell.font = Font(bold=True)
        for a in areas:
            row = [a]
            for cl in classes:
                cell = cell_map.get((a, cl))
                if not cell:
                    row.append("")
                else:
                    d = cell.get("last_depth") or "planned"
                    v = cell.get("visit_count") or 0
                    row.append(f"{d}x{v}" if v else str(d))
            cov_ws.append(row)

    if opts.include_codemap and ctx.get("codemap"):
        cm = ctx["codemap"]
        cm_ws = wb.create_sheet("Codemap")
        s = cm.get("summary") or {}
        cm_ws.append(["key", "value"])
        for k, v in (s if isinstance(s, dict) else {}).items():
            cm_ws.append([k, json.dumps(v) if isinstance(v, (dict, list)) else v])
        cm_ws.append([])
        cm_ws.append(["module_path", "kind", "file_count", "symbol_count"])
        for m in cm.get("modules") or []:
            if isinstance(m, dict):
                cm_ws.append(
                    [
                        m.get("path"),
                        m.get("kind"),
                        m.get("file_count"),
                        m.get("symbol_count"),
                    ]
                )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_add_md_lines(doc: Any, md_lines: list[str]) -> None:
    """Map simple markdown section lines into python-docx paragraphs."""
    for raw in md_lines:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.startswith(">"):
            doc.add_paragraph(line.lstrip("> ").strip())
        elif line.startswith("| "):
            doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)


def export_docx(
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> bytes:
    """Word document (.docx). Requires python-docx. Simple headings + bullets."""
    try:
        from docx import Document
    except ImportError as e:
        raise ImportError(
            "export_docx requires python-docx. Install with: pip install 'vulnforge[export]' "
            "or: pip install python-docx>=1.1"
        ) from e

    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    ctx = collect_report_context(run_dir, db, opts)
    rows = ctx.get("rows") or []
    counts = ctx.get("counts") or _summary_counts(rows)

    doc = Document()
    doc.add_heading("VulnForge audit report", level=0)
    doc.add_paragraph(f"Target: {ctx.get('target_path') or '-'}")
    doc.add_paragraph(f"Run: {ctx.get('run_id') or '-'}")
    doc.add_paragraph(f"Profile: {ctx.get('profile') or '-'}")
    doc.add_paragraph(f"Run dir: {ctx.get('run_dir') or '-'}")
    doc.add_paragraph(f"Exported: {ctx.get('exported_at') or utc_now_iso()}")
    sections_label = ", ".join(
        k.replace("include_", "")
        for k, v in sorted(opts.as_dict().items())
        if v and k.startswith("include_")
    ) or "none"
    doc.add_paragraph(f"Sections: {sections_label}")
    doc.add_paragraph(f"Include PoC: {'yes' if opts.include_poc else 'no'}")
    doc.add_paragraph(f"Disclaimer: {DISCLAIMER}")
    doc.add_paragraph(
        f"Needs human: {counts.get('needs_human', 0)} | "
        f"Confirmed (human): {counts.get('confirmed', 0)} | "
        f"Candidates: {counts.get('candidate', 0)} | "
        f"Rejected: {counts.get('rejected', 0)} | "
        f"Total: {counts.get('total', 0)}"
    )

    extra_md: list[str] = []
    if opts.include_summary:
        extra_md.extend(_md_summary_section(ctx.get("campaign"), counts))
    if opts.include_architecture:
        extra_md.extend(_md_architecture_section(ctx.get("architecture")))
    if opts.include_codemap:
        extra_md.extend(_md_codemap_section(ctx.get("codemap")))
    if opts.include_hunts:
        extra_md.extend(
            _md_coverage_section(ctx.get("coverage"), ctx.get("hunt_tasks") or [])
        )
    if extra_md:
        _docx_add_md_lines(doc, extra_md)

    if opts.include_findings:
        if not rows:
            doc.add_heading("Findings", level=1)
            doc.add_paragraph("No findings.")
        else:
            buckets = _partition_rows(rows)
            sections = [
                ("needs_human", "Needs human review (mech-passed)"),
                ("confirmed", "Confirmed findings (human-accepted)"),
                ("candidate", "Candidates"),
                ("rejected", "Rejected"),
                ("other", "Other"),
            ]
            for key, heading in sections:
                group = buckets.get(key) or []
                if key == "other" and not group:
                    continue
                doc.add_heading(heading, level=1)
                if not group:
                    doc.add_paragraph("None.")
                    continue
                for r in group:
                    body = r.get("_body") or {}
                    hr = r.get("_human_review") or {}
                    doc.add_heading(f"[{r['id']}] {r['title']}", level=2)
                    bullets = [
                        f"state: {r['state']}",
                        f"class: {r['weakness_class'] or '-'}",
                        f"severity: {r['severity'] or '-'}",
                        f"stable_key: {r['stable_key'] or '-'}",
                    ]
                    if r.get("attacker"):
                        bullets.append(f"attacker: {r['attacker']}")
                    if r.get("boundary"):
                        bullets.append(f"boundary: {r['boundary']}")
                    if r.get("impact"):
                        bullets.append(f"impact: {r['impact']}")
                    if r.get("path"):
                        loc = _path_label(
                            {
                                "path": r["path"],
                                "start_line": r["start_line"]
                                if r["start_line"] != ""
                                else None,
                                "end_line": r["end_line"]
                                if r["end_line"] != ""
                                else None,
                            }
                        )
                        bullets.append(f"location: {loc}")
                    if r.get("evidence_id"):
                        bullets.append(f"evidence: {r['evidence_id']}")
                    if r.get("poc_relpath"):
                        bullets.append(f"poc_relpath: {r['poc_relpath']}")
                    if hr.get("action"):
                        notes = hr.get("notes") or ""
                        note_bit = f" — {notes}" if notes else ""
                        bullets.append(f"human_review: {hr.get('action')}{note_bit}")
                    if r.get("validation_reasons"):
                        bullets.append(
                            f"validation_reasons: {r['validation_reasons']}"
                        )
                    for b in bullets:
                        doc.add_paragraph(b, style="List Bullet")
                    if r.get("summary"):
                        doc.add_paragraph(r["summary"])
                    cites = body.get("citations") or []
                    if cites:
                        doc.add_paragraph("Citations:")
                        for cit in cites:
                            if not isinstance(cit, dict) or not cit.get("path"):
                                continue
                            label = _path_label(cit)
                            if cit.get("symbol"):
                                label = f"{label} ({cit['symbol']})"
                            doc.add_paragraph(label, style="List Bullet")
                    pack = r.get("_poc") if opts.include_poc else None
                    if isinstance(pack, dict) and pack.get("files"):
                        doc.add_paragraph("PoC / evidence pack:")
                        for f in pack["files"]:
                            rel = f.get("relpath") or "?"
                            trunc = " (truncated)" if f.get("truncated") else ""
                            doc.add_heading(f"{rel}{trunc}", level=3)
                            doc.add_paragraph(str(f.get("content") or ""))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

# Map format query params to (exporter, extension, media_type, is_binary)
EXPORT_FORMATS: dict[str, tuple[str, str, str, bool]] = {
    "json": ("export_json", "json", "application/json; charset=utf-8", False),
    "md": ("export_markdown", "md", "text/markdown; charset=utf-8", False),
    "markdown": ("export_markdown", "md", "text/markdown; charset=utf-8", False),
    "csv": ("export_csv", "csv", "text/csv; charset=utf-8", False),
    "html": ("export_html", "html", "text/html; charset=utf-8", False),
    "xlsx": ("export_xlsx", "xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", True),
    "docx": (
        "export_docx",
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        True,
    ),
}


def normalize_formats(fmt: str | Iterable[str] | None) -> list[str]:
    """Parse format query (single, comma-separated, or list) → unique canonical keys."""
    if fmt is None:
        return ["json"]
    if isinstance(fmt, str):
        parts = [p.strip().lower() for p in fmt.replace(";", ",").split(",") if p.strip()]
    else:
        parts = [str(p).strip().lower() for p in fmt if str(p).strip()]
    seen: list[str] = []
    for p in parts:
        key = "md" if p == "markdown" else p
        if key not in EXPORT_FORMATS:
            raise ValueError(
                f"unknown export format {p!r}; expected one of: "
                + ", ".join(sorted({k for k in EXPORT_FORMATS if k != "markdown"}))
            )
        if key not in seen:
            seen.append(key)
    if not seen:
        raise ValueError("at least one export format is required")
    return seen

def export_bytes(
    fmt: str,
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> tuple[bytes, str, str]:
    """
    Dispatch export by format name.

    Returns (payload_bytes, filename_extension, media_type).
    Raises ValueError for unknown format; ImportError for missing optional deps.

    For multiple formats, use ``export_bundle`` (or pass comma-separated fmt via
    the API which zips automatically).
    """
    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    key = (fmt or "").strip().lower()
    if "," in key or ";" in key:
        return export_bundle(key, run_dir, db, options=opts)
    if key not in EXPORT_FORMATS:
        raise ValueError(
            f"unknown export format {fmt!r}; expected one of: "
            + ", ".join(sorted({k for k in EXPORT_FORMATS if k != "markdown"}))
        )
    fn_name, ext, media, is_binary = EXPORT_FORMATS[key]
    fn = globals()[fn_name]
    result = fn(run_dir, db, options=opts)
    if is_binary:
        if not isinstance(result, (bytes, bytearray)):
            raise TypeError(f"{fn_name} must return bytes")
        return bytes(result), ext, media
    if isinstance(result, bytes):
        return result, ext, media
    return str(result).encode("utf-8"), ext, media


def export_bundle(
    formats: str | Iterable[str],
    run_dir: Path | str | None,
    db: Any,
    *,
    include_poc: bool = False,
    options: ExportOptions | None = None,
    **section_flags: Any,
) -> tuple[bytes, str, str]:
    """
    Export one or more formats. Single format returns that file; multi → zip.

    Returns (payload_bytes, extension, media_type).
    """
    opts = ExportOptions.from_kwargs(
        include_poc=include_poc, options=options, **section_flags
    )
    keys = normalize_formats(formats)
    if len(keys) == 1:
        return export_bytes(keys[0], run_dir, db, options=opts)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for key in keys:
            payload, ext, _media = export_bytes(key, run_dir, db, options=opts)
            zf.writestr(f"vulnforge-report.{ext}", payload)
    return (
        buf.getvalue(),
        "zip",
        "application/zip",
    )


def export_filename(
    target_id: str,
    run_id: str,
    ext: str,
    *,
    include_poc: bool = False,
) -> str:
    """Safe Content-Disposition filename."""
    safe_t = "".join(c if c.isalnum() or c in "._-" else "-" for c in (target_id or "run"))
    safe_r = "".join(c if c.isalnum() or c in "._-" else "-" for c in (run_id or "export"))
    suffix = "-with-poc" if include_poc else ""
    return f"vulnforge-{safe_t}-{safe_r}-report{suffix}.{ext}"
